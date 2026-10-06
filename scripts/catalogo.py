#!/usr/bin/env python3
"""Motor do catálogo IT M&A (A&M · DTS).

Uso:
  python3 scripts/catalogo.py validar          # checa os dados e recalcula os índices
  python3 scripts/catalogo.py painel           # gera painel/index.html a partir de data/
  python3 scripts/catalogo.py extrair ARQ...   # extrai texto de .pptx/.pdf para fontes/brutos/

data/ofertas.yaml é a fonte única da verdade do portfólio atual; data/governanca.yaml é opcional;
data/portfolio-2027.csv (opcional) traz a proposta 2027 exibida no painel.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from pathlib import Path

import yaml

RAIZ = Path(__file__).resolve().parent.parent
DADOS = RAIZ / "data"
PAINEL_TEMPLATE = RAIZ / "painel" / "template.html"
PAINEL_SAIDA = RAIZ / "painel" / "index.html"
BRUTOS = RAIZ / "fontes" / "brutos"

CAMPOS_OBRIGATORIOS = (
    "id", "codigo", "nome", "familia", "fases", "readiness",
    "situacao_status_slide", "no_catalogo_cliente",
)


def carregar() -> tuple[dict, dict | None]:
    catalogo = yaml.safe_load((DADOS / "ofertas.yaml").read_text(encoding="utf-8"))
    arq_gov = DADOS / "governanca.yaml"
    governanca = yaml.safe_load(arq_gov.read_text(encoding="utf-8")) if arq_gov.exists() else None
    return catalogo, governanca


def nivel(valor: float) -> int:
    """Nível da escala oficial (1–5) em que um readiness se encontra."""
    return max(1, min(5, math.floor(valor + 1e-9)))


def media(valores: list[float]) -> float:
    return round(sum(valores) / len(valores), 2) if valores else 0.0


def validar(catalogo: dict, governanca: dict | None) -> list[str]:
    erros: list[str] = []
    ofertas = catalogo.get("ofertas", [])
    fases = {f["id"] for f in catalogo.get("fases", [])}
    familias = {f["id"] for f in catalogo.get("familias", [])}
    ids = [o.get("id") for o in ofertas]

    for dup in {i for i in ids if ids.count(i) > 1}:
        erros.append(f"id duplicado: {dup}")

    for o in ofertas:
        oid = o.get("id", "?")
        for campo in CAMPOS_OBRIGATORIOS:
            if campo not in o:
                erros.append(f"{oid}: campo obrigatório ausente '{campo}'")
        if o.get("familia") not in familias:
            erros.append(f"{oid}: família desconhecida '{o.get('familia')}'")
        for f in o.get("fases", []):
            if f not in fases:
                erros.append(f"{oid}: fase desconhecida '{f}'")
        r = o.get("readiness", {})
        atual, alvo = r.get("atual"), r.get("alvo_fy")
        if not all(isinstance(v, (int, float)) and 1 <= v <= 5 for v in (atual, alvo)):
            erros.append(f"{oid}: readiness fora da escala 1–5 ({atual} / {alvo})")
        elif alvo < atual:
            erros.append(f"{oid}: alvo FY ({alvo}) menor que o atual ({atual})")
        if o.get("situacao_status_slide") not in ("ativa", "riscada"):
            erros.append(f"{oid}: situacao_status_slide inválida")
        for c in (o.get("analise") or {}).get("conexoes", []):
            if c not in ids:
                erros.append(f"{oid}: conexão para oferta inexistente '{c}'")

    declarados = catalogo.get("indicadores_service_line", {})
    calc_atual = media([o["readiness"]["atual"] for o in ofertas])
    calc_alvo = media([o["readiness"]["alvo_fy"] for o in ofertas])
    if declarados.get("readiness_atual") != calc_atual:
        erros.append(f"readiness atual declarado {declarados.get('readiness_atual')} ≠ média calculada {calc_atual}")
    if declarados.get("readiness_alvo_fy") != calc_alvo:
        erros.append(f"readiness alvo declarado {declarados.get('readiness_alvo_fy')} ≠ média calculada {calc_alvo}")

    if governanca:
        for oid in governanca.get("ofertas", {}):
            if oid not in ids:
                erros.append(f"governanca.yaml: oferta inexistente '{oid}'")
        for oid in ids:
            if oid not in governanca.get("ofertas", {}):
                erros.append(f"governanca.yaml: oferta sem PO/squad '{oid}'")
    return erros


def relatorio(catalogo: dict, governanca: dict | None) -> str:
    escala = {e["nivel"]: e["nome"] for e in catalogo["escala_readiness"]}
    ofertas = catalogo["ofertas"]
    gov = (governanca or {}).get("ofertas", {})
    linhas = [
        "| Código | Oferta | Atual | Alvo FY | Gap | Nível atual | PO |",
        "|---|---|---:|---:|---:|---|---|",
    ]
    for o in sorted(ofertas, key=lambda x: -x["readiness"]["atual"]):
        r = o["readiness"]
        tachada = " ~~(tachada)~~" if o["situacao_status_slide"] == "riscada" else ""
        linhas.append(
            f"| {o['codigo']} | {o['nome']}{tachada} | {r['atual']:.1f} | {r['alvo_fy']:.1f} "
            f"| +{r['alvo_fy'] - r['atual']:.1f} | {escala[nivel(r['atual'])]} | {gov.get(o['id'], {}).get('po', '—')} |"
        )
    ativas = [o for o in ofertas if o["situacao_status_slide"] == "ativa"]
    linhas += [
        "",
        f"Readiness SL (10 linhas): atual {media([o['readiness']['atual'] for o in ofertas])} · "
        f"alvo {media([o['readiness']['alvo_fy'] for o in ofertas])}",
        f"Readiness SL (só linhas não tachadas, n={len(ativas)}): atual "
        f"{media([o['readiness']['atual'] for o in ativas])} · alvo {media([o['readiness']['alvo_fy'] for o in ativas])}",
        f"Ofertas em nível ≥ 4 (estruturada): hoje {sum(o['readiness']['atual'] >= 4 for o in ofertas)} · "
        f"no alvo FY {sum(o['readiness']['alvo_fy'] >= 4 for o in ofertas)}",
    ]
    return "\n".join(linhas)


def carregar_portfolio_2027() -> tuple[dict, list]:
    """Lê data/portfolio-2027.csv: devolve {código atual: linha da Tabela A} e a Tabela B."""
    arq = DADOS / "portfolio-2027.csv"
    if not arq.exists():
        return {}, []
    with arq.open(encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.DictReader(f))
    por_codigo, novos = {}, []
    for ln in linhas:
        if ln["tabela"] == "A":
            for cod in ln["codigo_hoje"].split(";"):
                por_codigo[cod.strip()] = ln
        else:
            novos.append({k: ln[k] for k in (
                "id", "nome_2027", "decisao_ou_tipo", "linha_2027", "o_que_e", "como_com_ia",
                "receita", "tem_recorrencia", "para_quem", "nota_prioridade_0_30", "onda",
            )})
    novos.sort(key=lambda n: -int(n["nota_prioridade_0_30"] or 0))
    return por_codigo, novos


def montar_dados_painel(catalogo: dict, governanca: dict | None) -> dict:
    gov = (governanca or {}).get("ofertas", {})
    prop_2027, novos_2027 = carregar_portfolio_2027()
    ofertas = []
    for o in catalogo["ofertas"]:
        item = {k: o.get(k) for k in (
            "id", "codigo", "nome", "nome_no_status", "nome_no_catalogo", "familia", "fases",
            "lado", "tipo_transacao", "situacao_status_slide", "no_catalogo_cliente",
            "readiness", "descricao_oficial", "dossie", "deck",
        )}
        item["proposta_de_valor"] = (o.get("analise") or {}).get("proposta_de_valor")
        if o["id"] in gov:
            item["po"] = gov[o["id"]].get("po")
            item["squad"] = gov[o["id"]].get("squad", [])
        if o["codigo"] in prop_2027:
            ln = prop_2027[o["codigo"]]
            item["proposta_2027"] = {k: ln[k] for k in (
                "id", "nome_2027", "decisao_ou_tipo", "linha_2027", "como_sem_ia", "como_com_ia",
                "receita", "onda",
            )}
        ofertas.append(item)
    return {
        "meta": {k: catalogo["meta"][k] for k in ("firma", "pratica", "service_line", "atualizado_em")},
        "escala": catalogo["escala_readiness"],
        "indicadores": catalogo["indicadores_service_line"],
        "fases": catalogo["fases"],
        "familias": catalogo["familias"],
        "lideranca": (governanca or {}).get("lideranca_service_line", []),
        "aviso": (governanca or {}).get("aviso_do_slide"),
        "ofertas": ofertas,
        "novos_2027": novos_2027,
    }


def gerar_painel(catalogo: dict, governanca: dict | None) -> Path:
    dados = montar_dados_painel(catalogo, governanca)
    # "</" escapado para que nenhum texto feche o <script> do painel.
    bloco = json.dumps(dados, ensure_ascii=False, indent=1).replace("</", "<\\/")
    html = PAINEL_TEMPLATE.read_text(encoding="utf-8").replace("/*__DADOS__*/{}", bloco)
    PAINEL_SAIDA.write_text(html, encoding="utf-8")
    return PAINEL_SAIDA


def slug(texto: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", texto.lower()).strip("-")[:60] or "arquivo"


def extrair(arquivos: list[str]) -> None:
    """Extrai texto por slide/página para fontes/brutos/ (pasta fora do git)."""
    for caminho in map(Path, arquivos):
        destino = BRUTOS / slug(caminho.stem)
        destino.mkdir(parents=True, exist_ok=True)
        if caminho.suffix.lower() == ".pptx":
            from pptx import Presentation  # pip install python-pptx

            paginas = []
            for slide in Presentation(str(caminho)).slides:
                textos = [
                    p.text.strip()
                    for forma in slide.shapes if forma.has_text_frame
                    for p in forma.text_frame.paragraphs if p.text.strip()
                ]
                notas = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
                paginas.append("\n".join(textos) + (f"\n\n> Notas: {notas}" if notas else ""))
        elif caminho.suffix.lower() == ".pdf":
            from pypdf import PdfReader  # pip install pypdf

            paginas = [(p.extract_text() or "").strip() for p in PdfReader(str(caminho)).pages]
        else:
            sys.exit(f"Formato não suportado: {caminho.suffix}")
        for n, texto in enumerate(paginas, 1):
            (destino / f"{n:03d}.md").write_text(f"# {caminho.name} · {n}\n\n{texto}\n", encoding="utf-8")
        print(f"{caminho.name}: {len(paginas)} páginas → {destino.relative_to(RAIZ)}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="comando", required=True)
    sub.add_parser("validar", help="checa os dados e imprime o relatório de readiness")
    sub.add_parser("painel", help="gera painel/index.html")
    p_ext = sub.add_parser("extrair", help="extrai texto de .pptx/.pdf para fontes/brutos/")
    p_ext.add_argument("arquivos", nargs="+")
    args = parser.parse_args()

    if args.comando == "extrair":
        extrair(args.arquivos)
        return 0

    catalogo, governanca = carregar()
    erros = validar(catalogo, governanca)
    if erros:
        print("ERROS DE VALIDAÇÃO:", *erros, sep="\n  - ")
        return 1
    if args.comando == "validar":
        print(relatorio(catalogo, governanca))
        print("\nOK: dados consistentes.")
    else:
        print(f"Painel gerado em {gerar_painel(catalogo, governanca).relative_to(RAIZ)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
